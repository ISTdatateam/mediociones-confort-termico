"""Shared helpers for building Word tables."""
from typing import Iterable, Sequence

from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor

from .report_data import ftemp


def set_column_width(table, col_index: int, width: Cm) -> None:
    if table is None:
        return
    for cell in table.columns[col_index].cells:
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        tcW = OxmlElement("w:tcW")
        tcW.set(qn("w:w"), str(int(width.inches * 1440)))
        tcW.set(qn("w:type"), "dxa")
        tcPr.append(tcW)


def merge_column_cells(table, col_index: int, start_row: int = 1) -> None:
    num_rows = len(table.rows)
    current_value = None
    start_index = None
    for i in range(start_row, num_rows):
        cell = table.cell(i, col_index)
        text = cell.text.strip()
        if text:
            if current_value is not None and start_index is not None and i - start_index > 1:
                first_cell = table.cell(start_index, col_index)
                for j in range(start_index + 1, i):
                    first_cell = first_cell.merge(table.cell(j, col_index))
            current_value = text
            start_index = i
    if current_value is not None and start_index is not None and num_rows - start_index > 1:
        first_cell = table.cell(start_index, col_index)
        for j in range(start_index + 1, num_rows):
            first_cell = first_cell.merge(table.cell(j, col_index))


def set_row_bold(row) -> None:
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True


def apply_header_style(cell, shading_color: str = "4f0b7b") -> None:
    shading = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls("w"), shading_color))
    cell._tc.get_or_add_tcPr().append(shading)
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True
            run.font.size = Pt(10)


def format_row(row, shading_color: str = "4f0b7b") -> None:
    for cell in row.cells:
        apply_header_style(cell, shading_color)


def add_row(table, label: str, value: str = "", first: bool = False):
    row = table.add_row()
    cells = row.cells
    if value == "":
        merged_cell = cells[0].merge(cells[1])
        merged_cell.text = label
        apply_header_style(merged_cell)
    else:
        cells[0].text = label
        cells[1].text = "" if (value is None or (isinstance(value, str) and value.strip() == "")) else str(value)
        if first:
            format_row(row)
    return row


def add_table_with_rows(doc, title: str, rows: Iterable[Sequence[str]]):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    add_row(table, title)
    for label, value in rows:
        add_row(table, label, value)
    return table


def set_bold_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    paragraph.text = ""
    run = paragraph.add_run(str(text))
    run.bold = True


def _format_metric(value, decimals: int) -> str:
    if value in (None, ""):
        return ""
    try:
        number = float(value)
    except (TypeError, ValueError):
        return ""
    return ftemp(f"{number:.{decimals}f}")


def add_summary_row(table, area: str, analisis: str, puesto_trabajo: str, metrics: dict) -> None:
    row_cells = table.add_row().cells
    set_bold_cell_text(row_cells[0], area)
    set_bold_cell_text(row_cells[1], analisis.upper())
    row_cells[2].text = puesto_trabajo
    metric_spec = [
        ("t_bul", 1),
        ("t_globo", 1),
        ("hum", 1),
        ("vel", 2),
        ("ppd", 1),
        ("pmv", 2),
    ]
    for index, (key, decimals) in enumerate(metric_spec, start=3):
        row_cells[index].text = _format_metric(metrics.get(key), decimals)


def agregar_contenido(cell, items) -> None:
    for item in items:
        cell.add_paragraph(item)
    if cell.paragraphs:
        first_paragraph = cell.paragraphs[0]
        p_element = first_paragraph._element
        p_element.getparent().remove(p_element)


def agregar_medidas_por_accion(table, medidas):
    for medida in medidas:
        for accion in medida["acciones"]:
            row_cells = table.add_row().cells
            agregar_contenido(row_cells[0], medida["areas"])
            agregar_contenido(row_cells[1], [accion])
            row_cells[2].text = medida["plazo"]
            for cell in row_cells:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    return table


def crear_tabla_recomendaciones(doc, tipo_medida: str, medidas):
    headers = ["Área", "Prescripción de medidas", "Plazo"]
    col_widths = [Cm(2), Cm(14), Cm(4)]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for idx, width in enumerate(col_widths):
        set_column_width(table, idx, width)
    header_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        header_cells[idx].text = header
    format_row(table.rows[0], shading_color="4F0B7B")
    agregar_medidas_por_accion(table, medidas)
    return table


__all__ = [
    "add_row",
    "add_summary_row",
    "add_table_with_rows",
    "agregar_contenido",
    "agregar_medidas_por_accion",
    "crear_tabla_recomendaciones",
    "format_row",
    "merge_column_cells",
    "set_bold_cell_text",
    "set_column_width",
    "set_row_bold",
]
