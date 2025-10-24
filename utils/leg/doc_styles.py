"""Utilities for configuring Word document styles and layout."""
import os
from typing import Dict

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def set_style_language(style, lang_code: str = "es-CL") -> None:
    rPr = style.element.get_or_add_rPr()
    lang = rPr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rPr.append(lang)
    lang.set(qn("w:val"), lang_code)
    lang.set(qn("w:eastAsia"), lang_code)
    lang.set(qn("w:bidi"), lang_code)


def apply_style_properties(style, properties: Dict) -> None:
    font = style.font
    font.name = properties.get("font_name", font.name)
    font.size = properties.get("font_size", font.size)
    font.color.rgb = properties.get("font_color", font.color.rgb)
    font.bold = properties.get("bold", font.bold)
    font.italic = properties.get("italic", font.italic)

    paragraph_format = style.paragraph_format
    paragraph_format.space_before = properties.get("space_before", paragraph_format.space_before)
    paragraph_format.space_after = properties.get("space_after", paragraph_format.space_after)
    paragraph_format.alignment = properties.get("alignment", paragraph_format.alignment)

    set_style_language(style, properties.get("lang_code", "es-CL"))


style_configurations = {
    "Normal": {
        "font_name": "Calibri",
        "font_size": Pt(10),
        "font_color": RGBColor(0x00, 0x00, 0x00),
        "bold": False,
        "italic": False,
        "space_before": Pt(0),
        "space_after": Pt(0),
        "alignment": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "lang_code": "es-CL",
    },
    "Heading 1": {
        "font_name": "Calibri",
        "font_size": Pt(14),
        "font_color": RGBColor(0, 0, 0),
        "bold": True,
        "italic": False,
        "space_before": Pt(0),
        "space_after": Pt(12),
        "alignment": WD_ALIGN_PARAGRAPH.LEFT,
        "lang_code": "es-CL",
    },
    "Heading 2": {
        "font_name": "Calibri",
        "font_size": Pt(12),
        "font_color": RGBColor(79, 11, 123),
        "bold": True,
        "italic": False,
        "space_before": Pt(12),
        "space_after": Pt(12),
        "alignment": WD_ALIGN_PARAGRAPH.LEFT,
        "lang_code": "es-CL",
    },
    "Heading 3": {
        "font_name": "Calibri",
        "font_size": Pt(11),
        "font_color": RGBColor(0x30, 0x30, 0x30),
        "bold": True,
        "italic": False,
        "space_before": Pt(8),
        "space_after": Pt(8),
        "alignment": WD_ALIGN_PARAGRAPH.LEFT,
        "lang_code": "es-CL",
    },
    "TablaTexto": {
        "font_name": "Calibri",
        "font_size": Pt(10),
        "font_color": RGBColor(0, 0, 0),
        "bold": False,
        "italic": False,
        "space_before": Pt(0),
        "space_after": Pt(0),
        "alignment": WD_ALIGN_PARAGRAPH.LEFT,
        "lang_code": "es-CL",
    },
    "Centrado": {
        "font_name": "Calibri",
        "font_size": Pt(10),
        "font_color": RGBColor(0, 0, 0),
        "bold": False,
        "italic": False,
        "space_before": Pt(0),
        "space_after": Pt(0),
        "alignment": WD_ALIGN_PARAGRAPH.CENTER,
        "lang_code": "es-CL",
    },
    "Centrado Bold": {
        "font_name": "Calibri",
        "font_size": Pt(10),
        "font_color": RGBColor(0, 0, 0),
        "bold": True,
        "italic": False,
        "space_before": Pt(0),
        "space_after": Pt(0),
        "alignment": WD_ALIGN_PARAGRAPH.CENTER,
        "lang_code": "es-CL",
    },
}


def look_informe(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2)
    for style_name, properties in style_configurations.items():
        try:
            style = doc.styles[style_name]
        except KeyError:
            continue
        apply_style_properties(style, properties)


def configurar_encabezado(doc: Document, logo_path: str = "assets/IST.jpg") -> None:
    section = doc.sections[0]
    section.header_distance = Cm(1.016)
    header = section.header
    paragraph = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    if os.path.exists(logo_path):
        run.add_picture(logo_path, width=Cm(2))


def agregar_titulo_y_codigo(doc: Document, titulo_texto: str, codigo_texto: str) -> None:
    titulo = doc.add_heading(titulo_texto, level=1)
    titulo.alignment = WD_ALIGN_PARAGRAPH.LEFT
    parrafo_codigo = doc.add_paragraph()
    run_codigo = parrafo_codigo.add_run(codigo_texto)
    run_codigo.bold = True
    parrafo_codigo.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def set_vertical_alignment(doc: Document, section_index: int = 0, alignment: str = "top") -> None:
    section = doc.sections[section_index]
    sectPr = section._sectPr
    vAlign = sectPr.find(qn("w:vAlign"))
    if vAlign is None:
        vAlign = OxmlElement("w:vAlign")
        sectPr.append(vAlign)
    alignment = alignment.lower()
    if alignment not in {"top", "center", "both"}:
        alignment = "top"
    vAlign.set(qn("w:val"), alignment)


__all__ = [
    "look_informe",
    "configurar_encabezado",
    "agregar_titulo_y_codigo",
    "set_vertical_alignment",
]
