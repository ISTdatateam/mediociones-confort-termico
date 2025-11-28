import os
import pandas as pd
from openpyxl import load_workbook
from copy import copy

# Carpeta donde están los archivos
carpeta = r"C:\Users\Quantum-Malloco\Desktop\Ventilacion\Casos de prueba 3"

# Texto a buscar (búsqueda parcial)
texto_buscar = "Identificación de aperturas por donde ingresa"
texto_normalizado = texto_buscar.strip().casefold()

# Mapeo de "Posición" (fila donde se encuentra el texto) -> (fila_en_la_que_inserto, cantidad_de_filas)
acciones = {
    50: (47, 28),
    57: (51, 21),
    55: (51, 23),
    54: (51, 24),
    53: (51, 25),
    58: (55, 20),
    62: (59, 16),
    60: (59, 18),
    66: (63, 12),
    71: (71, 7),
    42: (39, 36),
    45: (43, 33),
    49: (47, 29),
    51: (51, 31),
    52: (51, 30),
    59: (54, 19),
    61: (59, 18),
    65: (63, 13),
    67: (64, 11),
    70: (67, 8)
}

# Recorremos los archivos de la carpeta
lista_archivos = [
    f for f in os.listdir(carpeta)
    if os.path.isfile(os.path.join(carpeta, f))
]

for nombre_archivo in lista_archivos:
    # Solo procesar xlsx / xlsm
    if not nombre_archivo.lower().endswith(('.xlsx', '.xlsm')):
        continue

    ruta_archivo = os.path.join(carpeta, nombre_archivo)
    print(f"\nRevisando: {ruta_archivo}")

    # ---------------- BUSCAR POSICIÓN DEL TEXTO (con pandas) ----------------
    try:
        with pd.ExcelFile(ruta_archivo) as xls:
            primera_hoja = xls.sheet_names[0]

            df_cols = pd.read_excel(
                xls,
                sheet_name=primera_hoja,
                usecols="F:L",
                header=None
            )

    except ValueError:
        print("  La primera hoja no tiene columnas F:L. Se omite.")
        continue
    except Exception as e:
        print(f"  Error leyendo {nombre_archivo}: {e}")
        continue

    df_norm = (
        df_cols
        .fillna("")
        .astype(str)
        .apply(lambda col: col.str.strip().str.casefold())
    )

    df_mask = df_norm.apply(
        lambda col: col.str.contains(texto_normalizado, na=False)
    )

    mask_filas = df_mask.any(axis=1)
    indices_coinciden = mask_filas[mask_filas].index

    if len(indices_coinciden) == 0:
        print("  No se encontró el texto.")
        continue

    idx_primero = indices_coinciden[0]
    fila_excel = idx_primero + 1

    print(f"  >>> Texto encontrado en la fila {fila_excel}")

    # ---------------- DETERMINAR ACCIÓN ----------------
    if fila_excel not in acciones:
        print("  No hay acción configurada para esta posición.")
        continue

    fila_insertar, cantidad_filas = acciones[fila_excel]
    print(f"  Acción: insertar {cantidad_filas} filas en la fila {fila_insertar}")

    # ---------------- INSERTAR FILAS ----------------
    try:
        if nombre_archivo.lower().endswith('.xlsm'):
            wb = load_workbook(ruta_archivo, keep_vba=True)
        else:
            wb = load_workbook(ruta_archivo)

        ws = wb[primera_hoja]

        # Insertar filas
        ws.insert_rows(fila_insertar, amount=cantidad_filas)

        # ============================================================
        # >>> AQUI SE APLICA LA ESTRATEGIA SOLICITADA <<<
        # ============================================================

        inicio = fila_insertar
        fin = fila_insertar + cantidad_filas - 1

        # 1) Limpiar contenido y estilo de las filas nuevas
        for row in range(inicio, fin + 1):
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=row, column=col)
                cell.value = None
                cell._style = cell._style.__class__()  # estilo limpio

        # 2) Romper SOLO merges completamente dentro del bloque insertado
        rangos_merged = list(ws.merged_cells.ranges)

        for rango in rangos_merged:
            # Solo descombinar si TODO el rango está dentro de las filas nuevas
            if rango.min_row >= inicio and rango.max_row <= fin:
                ws.unmerge_cells(str(rango))

        # ============================================================

        wb.save(ruta_archivo)
        print(f"  Filas insertadas y limpiadas correctamente en {nombre_archivo}.")

    except Exception as e:
        print(f"  Error insertando filas en {nombre_archivo}: {e}")
        continue

print("\nProceso finalizado.")
